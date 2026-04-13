from pathlib import Path

from docx import Document

from app.models.schemas import ATSResume


class DocxExporter:
    def export(self, ats_resume: ATSResume, output_path: Path) -> Path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc = Document()

        name = ats_resume.header.get("name", "Candidate")
        doc.add_heading(name, level=0)
        for key in ("email", "phone", "location", "linkedin"):
            value = ats_resume.header.get(key)
            if value:
                doc.add_paragraph(f"{key.title()}: {value}")

        doc.add_heading("Summary", level=1)
        doc.add_paragraph(ats_resume.summary)

        doc.add_heading("Skills", level=1)
        doc.add_paragraph(", ".join(ats_resume.skills) if ats_resume.skills else "N/A")

        doc.add_heading("Experience", level=1)
        for exp in ats_resume.experience:
            title = exp.get("title", "")
            company = exp.get("company", "")
            doc.add_paragraph(f"{title} - {company}")
            for bullet in exp.get("bullets", []):
                doc.add_paragraph(str(bullet), style="List Bullet")

        doc.add_heading("Education", level=1)
        for edu in ats_resume.education:
            doc.add_paragraph(f"{edu.get('degree', '')} - {edu.get('institution', '')}")

        doc.add_heading("Certifications", level=1)
        for cert in ats_resume.certifications:
            doc.add_paragraph(cert, style="List Bullet")

        doc.save(output_path)
        return output_path
